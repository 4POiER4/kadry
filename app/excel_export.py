"""Формирование итогового .xlsx с подсветкой строк."""
from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FILLS = {
    "green": PatternFill("solid", fgColor="C6EFCE"),
    "red": PatternFill("solid", fgColor="FFC7CE"),
    "nodata": PatternFill("solid", fgColor="FFEB9C"),
    "onwork": PatternFill("solid", fgColor="DBEAFE"),
}
FONTS = {
    "green": Font(color="006100"),
    "red": Font(color="9C0006"),
    "nodata": Font(color="9C6500"),
    "onwork": Font(color="1E40AF"),
}
STATUS_TEXT = {
    "green": "Норма",
    "red": "Нарушение",
    "nodata": "Нет данных",
    "onwork": "На работе",
}

HEADERS = [
    "Таб. №", "Сотрудник",
    "Вход", "Выход", "Уйти не раньше", "Обед", "Зачёт (−обед)",
    "Норма (зачёт)", "Отклонение", "Статус", "Причина",
]


def build_xlsx(payload: dict[str, Any]) -> bytes:
    meta = payload.get("meta", {})
    rows = payload.get("rows", [])
    counts = meta.get("counts", {})

    wb = Workbook()
    ws = wb.active
    ws.title = "Присутствие"

    title = "Отчёт по присутствию"
    if meta.get("date"):
        title += f" за {meta['date']}"
    if meta.get("weekday_name"):
        title += f" ({meta['weekday_name']})"
    ws.append([title])
    ws.append([
        f"Режим: {meta.get('day_type_label', '')} | "
        f"норма присутствия {meta.get('required_presence', '')} | "
        f"обед {meta.get('lunch', '')} | "
        f"зачётная норма {meta.get('required_net', '')} | "
        f"окно прихода {meta.get('arrival_window', '')}"
    ])
    ws.append([
        f"Итого: {counts.get('total', 0)} | "
        f"Норма: {counts.get('green', 0)} | "
        f"На работе: {counts.get('onwork', 0)} | "
        f"Нарушения: {counts.get('red', 0)} | "
        f"Нет данных: {counts.get('nodata', 0)} | "
        f"сформировано {meta.get('generated_at', '')}"
    ])
    ws.append([])

    header_row = 5
    ws.append(HEADERS)
    for c in range(1, len(HEADERS) + 1):
        ws.cell(row=header_row, column=c).font = Font(bold=True)

    for r in rows:
        status = r.get("status", "nodata")
        countable = status not in ("nodata",)
        ws.append([
            r.get("tab", ""),
            r.get("name", ""),
            r.get("t_in", ""),
            r.get("t_out", ""),
            r.get("can_leave", ""),
            meta.get("lunch", "") if countable else "",
            r.get("net", ""),
            meta.get("required_net", "") if countable else "",
            r.get("delta", ""),
            STATUS_TEXT.get(status, status),
            r.get("reason", ""),
        ])
        row_i = ws.max_row
        fill = FILLS.get(status)
        font = FONTS.get(status)
        for c in range(1, len(HEADERS) + 1):
            cellobj = ws.cell(row=row_i, column=c)
            if fill:
                cellobj.fill = fill
            if font:
                cellobj.font = font

    widths = [10, 34, 8, 8, 13, 8, 13, 13, 12, 13, 46]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(HEADERS))}{ws.max_row}"
    for c in (3, 4, 5, 6, 7, 8, 9):
        for row_i in range(header_row + 1, ws.max_row + 1):
            ws.cell(row=row_i, column=c).alignment = Alignment(horizontal="center")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
