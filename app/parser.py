"""Чтение выгрузки из проходной системы и поиск структуры таблицы."""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
from typing import Any

WEEKDAYS_RU = [
    "понедельник", "вторник", "среда", "четверг", "пятница", "суббота", "воскресенье",
]


def _s(v: Any) -> str:
    return "" if v is None else str(v).strip()


def cell(row: list[Any], j: int | None) -> Any:
    if j is None or j < 0 or j >= len(row):
        return None
    return row[j]


# --------------------------------------------------------------------------- #
# Загрузка «сырых» строк из файла любого поддерживаемого формата
# --------------------------------------------------------------------------- #
def load_rows(filename: str, data: bytes) -> list[list[Any]]:
    name = (filename or "").lower()

    if name.endswith(".csv") or name.endswith(".txt"):
        text = data.decode("utf-8-sig", errors="replace")
        sample = text[:8192]
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        return [list(r) for r in reader]

    if name.endswith(".xls"):
        import xlrd  # type: ignore

        book = xlrd.open_workbook(file_contents=data)
        sheet = book.sheet_by_index(0)
        rows: list[list[Any]] = []
        for i in range(sheet.nrows):
            row: list[Any] = []
            for j in range(sheet.ncols):
                c = sheet.cell(i, j)
                if c.ctype == xlrd.XL_CELL_DATE:
                    y, mo, d, h, mi, s = xlrd.xldate_as_tuple(c.value, book.datemode)
                    if y == mo == d == 0:
                        row.append(dt.time(h, mi, s))
                    else:
                        row.append(dt.datetime(y, mo, d, h, mi, s))
                else:
                    row.append(c.value)
            rows.append(row)
        return rows

    # xlsx / xlsm по умолчанию
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb.active
    return [list(r) for r in ws.iter_rows(values_only=True)]


# --------------------------------------------------------------------------- #
# Поиск строки заголовков и раскладки колонок
# --------------------------------------------------------------------------- #
def find_layout(rows: list[list[Any]]) -> tuple[int | None, dict[str, int]]:
    for idx, row in enumerate(rows[:40]):
        low = [_s(c).lower() for c in row]
        has_tab = any("таб" in c for c in low)
        has_name = any("сотрудн" in c or "фамилия" in c or "ф.и.о" in c or "фио" in c for c in low)
        if not (has_tab and has_name):
            continue

        cols: dict[str, int] = {}
        for j, c in enumerate(low):
            if not c:
                continue
            if "таб" in c:
                cols.setdefault("tab", j)
            elif "сотрудн" in c or "фамилия" in c or "ф.и.о" in c or c == "фио":
                cols.setdefault("name", j)
            elif "подраздел" in c or "отдел" in c:
                cols.setdefault("dep", j)
            elif "должност" in c:
                cols.setdefault("pos", j)
            elif "вход" in c:
                cols.setdefault("t_in", j)
            elif "выход" in c:
                cols.setdefault("t_out", j)
            elif "присутств" in c:
                cols.setdefault("pres", j)
            elif c == "всего":
                cols.setdefault("total", j)

        cols.setdefault("t_in", 4)   # столбец E
        cols.setdefault("t_out", 5)  # столбец F
        cols.setdefault("name", 1)
        return idx, cols

    # Заголовки не нашли — работаем по позициям A..F
    return None, {"tab": 0, "name": 1, "dep": 2, "pos": 3, "t_in": 4, "t_out": 5}


# --------------------------------------------------------------------------- #
# Определение даты и дня недели отчёта
# --------------------------------------------------------------------------- #
def detect_day(rows: list[list[Any]], header_idx: int | None, cols: dict[str, int]) -> dict[str, Any]:
    chunks: list[str] = []
    if header_idx is not None:
        for key in ("t_in", "t_out", "pres"):
            chunks.append(_s(cell(rows[header_idx], cols.get(key))))
    for r in rows[: (header_idx + 1) if header_idx is not None else 8]:
        chunks.append(" ".join(_s(c) for c in r))
    text = " ".join(chunks).lower()

    date: dt.date | None = None
    m = re.search(r"(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{4})", text)
    if m:
        try:
            date = dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:
            date = None

    wd_idx: int | None = None
    for i, w in enumerate(WEEKDAYS_RU):
        if w[:5] in text:
            wd_idx = i
            break

    if date is not None:
        wd_idx = date.weekday()

    if wd_idx is None:
        return {"date": None, "weekday": None, "weekday_name": None, "day_type": None}

    day_type = "fri" if wd_idx == 4 else ("weekend" if wd_idx >= 5 else "week")
    return {
        "date": date.isoformat() if date else None,
        "weekday": wd_idx,
        "weekday_name": WEEKDAYS_RU[wd_idx],
        "day_type": day_type,
    }
